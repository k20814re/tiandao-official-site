#!/usr/bin/env python3
from __future__ import annotations
import json
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
BOOK=ROOT/'TianDao_B2B_Products.xlsx'
OUTPUT=ROOT/'data'/'b2b-products.json'
HEADERS={
 '商品':['啟用','排序','商品ID','圖示','中文名稱','英文名稱','分類ID','供應商ID','中文說明','英文說明','MOQ中文','MOQ英文','報價方式中文','報價方式英文','精選'],
 '分類':['啟用','排序','分類ID','圖示','中文名稱','英文名稱','中文說明','英文說明'],
 '供應商':['啟用','排序','供應商ID','圖示','中文名稱','英文名稱','分類ID','中文簡介','英文簡介','中文主要品項','英文主要品項'],
 '設定':['設定Key','值','用途']}
def clean(v:Any)->str:return '' if v is None else str(v).strip()
def yes(v:Any)->bool:return clean(v).lower() in {'是','true','1','yes','y'}
def num(v:Any)->int:
 try:return int(float(v))
 except (TypeError,ValueError):return 0
def rows(wb,name):
 if name not in wb.sheetnames:raise ValueError(f'找不到工作表：{name}')
 ws=wb[name];expected=HEADERS[name];actual=[clean(ws.cell(4,i).value) for i in range(1,len(expected)+1)]
 if actual!=expected:raise ValueError(f'{name} 第 4 列欄位不可更名。預期 {expected}，目前為 {actual}')
 return [dict(zip(expected,v)) for v in ws.iter_rows(min_row=5,max_col=len(expected),values_only=True) if any(clean(x) for x in v)]
def unique(items,label):
 ids=[x['id'] for x in items]
 if '' in ids:raise ValueError(f'{label} ID 不可空白')
 dup=sorted({x for x in ids if ids.count(x)>1})
 if dup:raise ValueError(f'{label} ID 重複：{dup}')
def build(wb):
 cats=[{'order':num(r['排序']),'id':clean(r['分類ID']),'icon':clean(r['圖示']),'nameZh':clean(r['中文名稱']),'nameEn':clean(r['英文名稱']),'descriptionZh':clean(r['中文說明']),'descriptionEn':clean(r['英文說明'])} for r in rows(wb,'分類') if yes(r['啟用'])]
 sups=[{'order':num(r['排序']),'id':clean(r['供應商ID']),'icon':clean(r['圖示']),'nameZh':clean(r['中文名稱']),'nameEn':clean(r['英文名稱']),'categoryId':clean(r['分類ID']),'descriptionZh':clean(r['中文簡介']),'descriptionEn':clean(r['英文簡介']),'itemsZh':clean(r['中文主要品項']),'itemsEn':clean(r['英文主要品項'])} for r in rows(wb,'供應商') if yes(r['啟用'])]
 prods=[{'order':num(r['排序']),'id':clean(r['商品ID']),'icon':clean(r['圖示']),'nameZh':clean(r['中文名稱']),'nameEn':clean(r['英文名稱']),'categoryId':clean(r['分類ID']),'supplierId':clean(r['供應商ID']),'descriptionZh':clean(r['中文說明']),'descriptionEn':clean(r['英文說明']),'moqZh':clean(r['MOQ中文']),'moqEn':clean(r['MOQ英文']),'pricingZh':clean(r['報價方式中文']),'pricingEn':clean(r['報價方式英文']),'featured':yes(r['精選'])} for r in rows(wb,'商品') if yes(r['啟用'])]
 for x,n in [(cats,'分類'),(sups,'供應商'),(prods,'商品')]:unique(x,n);x.sort(key=lambda a:a['order'])
 catids={x['id'] for x in cats};supids={x['id'] for x in sups}
 for x in sups:
  if x['categoryId'] not in catids:raise ValueError(f"供應商 {x['id']} 使用不存在的分類 {x['categoryId']}")
 for x in prods:
  if x['categoryId'] not in catids:raise ValueError(f"商品 {x['id']} 使用不存在的分類 {x['categoryId']}")
  if x['supplierId'] not in supids:raise ValueError(f"商品 {x['id']} 使用不存在的供應商 {x['supplierId']}")
 settings={clean(r['設定Key']):r['值'] for r in rows(wb,'設定') if clean(r['設定Key'])}
 return {'schemaVersion':num(settings.get('schemaVersion',1)),'settings':settings,'categories':cats,'suppliers':sups,'products':prods}
def main():
 if not BOOK.exists():raise FileNotFoundError(f'找不到 {BOOK.name}')
 wb=load_workbook(BOOK,data_only=True,read_only=True)
 try:data=build(wb)
 finally:wb.close()
 OUTPUT.parent.mkdir(parents=True,exist_ok=True);OUTPUT.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8');print(f'已更新 {OUTPUT.relative_to(ROOT)}：{len(data["products"])} 商品、{len(data["categories"])} 分類、{len(data["suppliers"])} 供應商')
if __name__=='__main__':main()
